from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor


# =========================================================
# CIVICEYE AI - PREMIUM YCS 2026 PRESENTATION
# 7 SLIDES
# =========================================================

prs = Presentation()

prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


# =========================================================
# COLORS
# =========================================================

BG = RGBColor(246, 248, 255)

DARK = RGBColor(25, 31, 52)
NAVY = RGBColor(20, 27, 51)

PURPLE = RGBColor(91, 76, 220)
PURPLE_DARK = RGBColor(70, 56, 180)
PURPLE_LIGHT = RGBColor(226, 222, 255)

BLUE = RGBColor(57, 168, 255)
BLUE_LIGHT = RGBColor(224, 241, 255)

CYAN = RGBColor(46, 201, 255)

GREEN = RGBColor(36, 180, 120)
GREEN_LIGHT = RGBColor(224, 248, 239)

ORANGE = RGBColor(245, 158, 11)
ORANGE_LIGHT = RGBColor(255, 244, 220)

WHITE = RGBColor(255, 255, 255)

GREY = RGBColor(100, 108, 128)
LIGHT = RGBColor(225, 230, 242)


# =========================================================
# BASIC HELPERS
# =========================================================

def add_text(
    slide,
    text,
    x,
    y,
    w,
    h,
    size=18,
    color=DARK,
    bold=False,
    align=PP_ALIGN.LEFT
):

    box = slide.shapes.add_textbox(
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h)
    )

    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True

    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0

    tf.vertical_anchor = MSO_ANCHOR.MIDDLE

    p = tf.paragraphs[0]

    p.text = text
    p.font.name = "Arial"
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = align

    return box


def add_background(slide, color=BG):

    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_card(
    slide,
    x,
    y,
    w,
    h,
    fill=WHITE,
    line=LIGHT,
    radius=True
):

    shape_type = (
        MSO_SHAPE.ROUNDED_RECTANGLE
        if radius
        else MSO_SHAPE.RECTANGLE
    )

    card = slide.shapes.add_shape(
        shape_type,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h)
    )

    card.fill.solid()
    card.fill.fore_color.rgb = fill

    card.line.color.rgb = line
    card.line.width = Pt(1)

    return card


def add_accent_line(slide, x, y, w, color=PURPLE):

    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(0.06)
    )

    line.fill.solid()
    line.fill.fore_color.rgb = color
    line.line.fill.background()

    return line


def add_header(slide, title, subtitle, number):

    add_text(
        slide,
        title,
        0.72,
        0.40,
        9.5,
        0.55,
        size=26,
        color=DARK,
        bold=True
    )

    add_text(
        slide,
        subtitle,
        0.75,
        1.00,
        10.4,
        0.4,
        size=13,
        color=GREY
    )

    add_text(
        slide,
        f"CIVICEYE AI   |   {number}/7",
        10.35,
        0.47,
        2.2,
        0.3,
        size=9,
        color=GREY,
        bold=True,
        align=PP_ALIGN.RIGHT
    )

    add_accent_line(
        slide,
        0.72,
        1.45,
        1.25,
        PURPLE
    )


def add_footer(slide):

    add_text(
        slide,
        "CivicEye AI  |  YCS 2026  |  Sri Lanka",
        0.72,
        7.12,
        5.0,
        0.2,
        size=8,
        color=GREY
    )


def add_pill(
    slide,
    text,
    x,
    y,
    w,
    fill,
    color=DARK
):

    add_card(
        slide,
        x,
        y,
        w,
        0.48,
        fill,
        fill
    )

    add_text(
        slide,
        text,
        x,
        y + 0.02,
        w,
        0.4,
        size=9,
        color=color,
        bold=True,
        align=PP_ALIGN.CENTER
    )


def add_arrow(
    slide,
    x1,
    y1,
    x2,
    y2,
    color=PURPLE
):

    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2)
    )

    connector.line.color.rgb = color
    connector.line.width = Pt(2)

    return connector


def add_feature_card(
    slide,
    title,
    description,
    x,
    y,
    w,
    h,
    fill,
    accent
):

    add_card(
        slide,
        x,
        y,
        w,
        h,
        fill,
        fill
    )

    add_accent_line(
        slide,
        x + 0.18,
        y + 0.18,
        0.45,
        accent
    )

    add_text(
        slide,
        title,
        x + 0.18,
        y + 0.38,
        w - 0.36,
        0.32,
        size=12,
        color=accent,
        bold=True
    )

    add_text(
        slide,
        description,
        x + 0.18,
        y + 0.76,
        w - 0.36,
        h - 0.88,
        size=10,
        color=DARK
    )


# =========================================================
# SLIDE 1
# COVER
# =========================================================

slide = prs.slides.add_slide(prs.slide_layouts[6])

add_background(slide, BG)


# Decorative background

circle = slide.shapes.add_shape(
    MSO_SHAPE.OVAL,
    Inches(9.25),
    Inches(-1.15),
    Inches(5.2),
    Inches(5.2)
)

circle.fill.solid()
circle.fill.fore_color.rgb = PURPLE_LIGHT
circle.line.fill.background()


circle2 = slide.shapes.add_shape(
    MSO_SHAPE.OVAL,
    Inches(-1.35),
    Inches(5.35),
    Inches(4.1),
    Inches(4.1)
)

circle2.fill.solid()
circle2.fill.fore_color.rgb = BLUE_LIGHT
circle2.line.fill.background()


# Logo

logo = add_card(
    slide,
    0.9,
    1.0,
    0.95,
    0.95,
    PURPLE,
    PURPLE
)

add_text(
    slide,
    "CE",
    0.9,
    1.15,
    0.95,
    0.55,
    size=21,
    color=WHITE,
    bold=True,
    align=PP_ALIGN.CENTER
)


# Title

add_text(
    slide,
    "CivicEye AI",
    2.05,
    1.02,
    6.7,
    0.75,
    size=40,
    color=DARK,
    bold=True
)

add_text(
    slide,
    "AI-Assisted Public Issue Reporting",
    2.08,
    1.88,
    7.4,
    0.55,
    size=21,
    color=PURPLE,
    bold=True
)

add_text(
    slide,
    "A digital platform for reporting, analysing, prioritising and monitoring public issues.",
    2.08,
    2.62,
    7.4,
    0.8,
    size=16,
    color=GREY
)


# Feature pills

add_pill(
    slide,
    "AI ANALYSIS",
    2.08,
    3.75,
    1.65,
    PURPLE_LIGHT,
    PURPLE
)

add_pill(
    slide,
    "LOCATION",
    3.92,
    3.75,
    1.45,
    BLUE_LIGHT,
    BLUE
)

add_pill(
    slide,
    "PRIORITY",
    5.55,
    3.75,
    1.45,
    GREEN_LIGHT,
    GREEN
)

add_pill(
    slide,
    "TRACKING",
    7.18,
    3.75,
    1.5,
    ORANGE_LIGHT,
    ORANGE
)


# Competition box

add_card(
    slide,
    2.08,
    4.75,
    6.95,
    0.95,
    PURPLE,
    PURPLE
)

add_text(
    slide,
    "Young Computer Scientist Competition (YCS) 2026",
    2.25,
    4.94,
    6.6,
    0.45,
    size=17,
    color=WHITE,
    bold=True,
    align=PP_ALIGN.CENTER
)


# =========================================================
# RIGHT SIDE - CIVICEYE MINI WORKFLOW
# =========================================================

add_card(
    slide,
    9.05,
    2.65,
    3.35,
    3.25,
    WHITE,
    LIGHT
)

add_text(
    slide,
    "CIVICEYE WORKFLOW",
    9.35,
    2.92,
    2.75,
    0.35,
    size=11,
    color=PURPLE,
    bold=True,
    align=PP_ALIGN.CENTER
)


mini_steps = [
    ("REPORT", PURPLE_LIGHT, PURPLE),
    ("AI ANALYSE", BLUE_LIGHT, BLUE),
    ("PRIORITY", GREEN_LIGHT, GREEN),
    ("AUTHORITY", ORANGE_LIGHT, ORANGE),
    ("TRACK", PURPLE_LIGHT, PURPLE)
]


for i, (label, fill, accent) in enumerate(mini_steps):

    y = 3.45 + i * 0.48

    add_card(
        slide,
        9.45,
        y,
        2.55,
        0.36,
        fill,
        fill
    )

    add_text(
        slide,
        label,
        9.45,
        y + 0.01,
        2.55,
        0.27,
        size=8.5,
        color=accent,
        bold=True,
        align=PP_ALIGN.CENTER
    )

    if i < len(mini_steps) - 1:

        add_arrow(
            slide,
            10.72,
            y + 0.36,
            10.72,
            y + 0.48,
            CYAN
        )


add_text(
    slide,
    "From citizen report to\ncontinuous monitoring",
    9.35,
    5.78,
    2.75,
    0.55,
    size=10,
    color=GREY,
    bold=True,
    align=PP_ALIGN.CENTER
)


add_text(
    slide,
    "Sri Lanka",
    9.05,
    6.45,
    3.35,
    0.35,
    size=13,
    color=PURPLE,
    bold=True,
    align=PP_ALIGN.CENTER
)


# =========================================================
# SLIDE 2
# THE PROBLEM
# =========================================================

slide = prs.slides.add_slide(prs.slide_layouts[6])

add_background(slide)

add_header(
    slide,
    "The Public Issue Challenge",
    "Common gaps in reporting, routing and monitoring public issues",
    2
)


problems = [
    (
        "01",
        "REPORTING",
        "Citizens need a structured way to submit issue details, images and location.",
        BLUE
    ),
    (
        "02",
        "VISIBILITY",
        "Citizens need clearer visibility after a report has been submitted.",
        PURPLE
    ),
    (
        "03",
        "DUPLICATES",
        "The same issue can potentially be reported multiple times.",
        GREEN
    ),
    (
        "04",
        "ROUTING",
        "Different issue categories may require different responsible authorities.",
        ORANGE
    )
]


positions = [
    (0.8, 1.95),
    (6.75, 1.95),
    (0.8, 4.25),
    (6.75, 4.25)
]


for (num, title, desc, accent), (x, y) in zip(
    problems,
    positions
):

    add_card(
        slide,
        x,
        y,
        5.75,
        1.8,
        WHITE,
        LIGHT
    )

    add_card(
        slide,
        x + 0.25,
        y + 0.27,
        0.62,
        0.62,
        accent,
        accent
    )

    add_text(
        slide,
        num,
        x + 0.25,
        y + 0.35,
        0.62,
        0.3,
        size=11,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER
    )

    add_text(
        slide,
        title,
        x + 1.05,
        y + 0.22,
        4.25,
        0.4,
        size=16,
        color=DARK,
        bold=True
    )

    add_text(
        slide,
        desc,
        x + 1.05,
        y + 0.75,
        4.25,
        0.7,
        size=12,
        color=GREY
    )


add_footer(slide)


# =========================================================
# SLIDE 3
# CIVICEYE SOLUTION + FEATURES
# =========================================================

slide = prs.slides.add_slide(prs.slide_layouts[6])

add_background(slide)

add_header(
    slide,
    "Introducing CivicEye AI",
    "A single digital workflow connecting citizens, analysis and administration",
    3
)


# Main statement

add_card(
    slide,
    0.75,
    1.72,
    12.0,
    0.95,
    PURPLE,
    PURPLE
)

add_text(
    slide,
    "Report  ->  Analyse  ->  Prioritise  ->  Route  ->  Monitor",
    1.05,
    1.95,
    11.4,
    0.45,
    size=19,
    color=WHITE,
    bold=True,
    align=PP_ALIGN.CENTER
)


# Feature cards

features = [
    (
        "IMAGE REPORTING",
        "Upload an image with a public issue report.",
        PURPLE_LIGHT,
        PURPLE
    ),
    (
        "LOCATION",
        "Capture the issue location using browser geolocation.",
        BLUE_LIGHT,
        BLUE
    ),
    (
        "VOICE",
        "Voice input supports English, Tamil and Sinhala.",
        GREEN_LIGHT,
        GREEN
    ),
    (
        "PRIORITY",
        "Generate an emergency priority score from report information.",
        ORANGE_LIGHT,
        ORANGE
    ),
    (
        "AUTHORITY",
        "Recommend a relevant authority based on issue category.",
        PURPLE_LIGHT,
        PURPLE
    ),
    (
        "DUPLICATES",
        "Check for potentially matching reports.",
        BLUE_LIGHT,
        BLUE
    ),
    (
        "NOTIFICATIONS",
        "Notify citizens when report information changes.",
        GREEN_LIGHT,
        GREEN
    ),
    (
        "TRACKING",
        "Follow report status from submission to resolution.",
        ORANGE_LIGHT,
        ORANGE
    )
]


for i, (title, desc, fill, accent) in enumerate(features):

    col = i % 4
    row = i // 4

    x = 0.75 + col * 3.05
    y = 3.00 + row * 1.65

    add_feature_card(
        slide,
        title,
        desc,
        x,
        y,
        2.75,
        1.35,
        fill,
        accent
    )


add_footer(slide)


# =========================================================
# SLIDE 4
# HOW IT WORKS
# =========================================================

slide = prs.slides.add_slide(prs.slide_layouts[6])

add_background(slide)

add_header(
    slide,
    "How CivicEye AI Works",
    "A complete workflow from citizen submission to status monitoring",
    4
)


steps = [
    (
        "01",
        "REPORT",
        "Citizen submits an image, description and location."
    ),
    (
        "02",
        "ANALYSE",
        "Rule-based logic analyses issue information."
    ),
    (
        "03",
        "PRIORITISE",
        "Priority and duplicate information are generated."
    ),
    (
        "04",
        "ROUTE",
        "A relevant authority is recommended."
    ),
    (
        "05",
        "REVIEW",
        "Admin reviews reports and updates status."
    ),
    (
        "06",
        "NOTIFY",
        "Citizen tracks the report and receives updates."
    )
]


for i, (num, title, desc) in enumerate(steps):

    col = i % 3
    row = i // 3

    x = 0.7 + col * 4.15
    y = 1.85 + row * 2.35

    accent = PURPLE if i % 2 == 0 else BLUE

    add_card(
        slide,
        x,
        y,
        3.7,
        1.8,
        WHITE,
        LIGHT
    )

    add_card(
        slide,
        x + 0.25,
        y + 0.25,
        0.65,
        0.65,
        accent,
        accent
    )

    add_text(
        slide,
        num,
        x + 0.25,
        y + 0.34,
        0.65,
        0.3,
        size=11,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER
    )

    add_text(
        slide,
        title,
        x + 1.1,
        y + 0.22,
        2.2,
        0.4,
        size=15,
        color=DARK,
        bold=True
    )

    add_text(
        slide,
        desc,
        x + 1.1,
        y + 0.72,
        2.25,
        0.7,
        size=10.5,
        color=GREY
    )


# Flow connectors

add_arrow(
    slide,
    4.4,
    2.75,
    4.75,
    2.75,
    CYAN
)

add_arrow(
    slide,
    8.55,
    2.75,
    8.9,
    2.75,
    CYAN
)

add_arrow(
    slide,
    4.4,
    5.1,
    4.75,
    5.1,
    CYAN
)

add_arrow(
    slide,
    8.55,
    5.1,
    8.9,
    5.1,
    CYAN
)


add_footer(slide)


# =========================================================
# SLIDE 5
# SYSTEM ARCHITECTURE
# =========================================================

slide = prs.slides.add_slide(prs.slide_layouts[6])

add_background(slide)

add_header(
    slide,
    "CivicEye AI System Architecture",
    "How the web application connects citizens, backend services and administration",
    5
)


# Left section

add_card(
    slide,
    0.7,
    1.75,
    3.65,
    4.95,
    WHITE,
    LIGHT
)

add_text(
    slide,
    "CITIZEN LAYER",
    1.0,
    2.05,
    3.05,
    0.4,
    size=16,
    color=PURPLE,
    bold=True,
    align=PP_ALIGN.CENTER
)


citizen_items = [
    "Report Form",
    "Image Upload",
    "Location",
    "Voice Input",
    "Report Tracking",
    "Notifications"
]

for i, item in enumerate(citizen_items):

    y = 2.72 + i * 0.55

    add_card(
        slide,
        1.05,
        y,
        2.95,
        0.38,
        PURPLE_LIGHT,
        PURPLE_LIGHT
    )

    add_text(
        slide,
        item,
        1.05,
        y + 0.02,
        2.95,
        0.3,
        size=10,
        color=PURPLE_DARK,
        bold=True,
        align=PP_ALIGN.CENTER
    )


# Middle section

add_card(
    slide,
    4.85,
    1.75,
    3.55,
    4.95,
    PURPLE,
    PURPLE
)

add_text(
    slide,
    "PROCESSING LAYER",
    5.15,
    2.05,
    2.95,
    0.4,
    size=16,
    color=WHITE,
    bold=True,
    align=PP_ALIGN.CENTER
)


processing = [
    "REST API",
    "Node.js / Express",
    "Image Handling",
    "Rule-Based Analysis",
    "Priority Logic",
    "Duplicate Check",
    "Authority Mapping"
]

for i, item in enumerate(processing):

    y = 2.65 + i * 0.5

    add_card(
        slide,
        5.25,
        y,
        2.75,
        0.35,
        WHITE,
        WHITE
    )

    add_text(
        slide,
        item,
        5.25,
        y + 0.01,
        2.75,
        0.28,
        size=9,
        color=DARK,
        bold=True,
        align=PP_ALIGN.CENTER
    )


# Right section

add_card(
    slide,
    8.9,
    1.75,
    3.7,
    4.95,
    WHITE,
    LIGHT
)

add_text(
    slide,
    "ADMIN LAYER",
    9.2,
    2.05,
    3.1,
    0.4,
    size=16,
    color=BLUE,
    bold=True,
    align=PP_ALIGN.CENTER
)


admin_items = [
    "Admin Dashboard",
    "Report Review",
    "Status Updates",
    "Notifications",
    "Resolution Verification",
    "Monitoring"
]

for i, item in enumerate(admin_items):

    y = 2.72 + i * 0.55

    add_card(
        slide,
        9.25,
        y,
        3.0,
        0.38,
        BLUE_LIGHT,
        BLUE_LIGHT
    )

    add_text(
        slide,
        item,
        9.25,
        y + 0.02,
        3.0,
        0.3,
        size=10,
        color=BLUE,
        bold=True,
        align=PP_ALIGN.CENTER
    )


# Architecture connectors

add_arrow(
    slide,
    4.35,
    4.1,
    4.8,
    4.1,
    CYAN
)

add_arrow(
    slide,
    8.45,
    4.1,
    8.85,
    4.1,
    CYAN
)


add_footer(slide)


# =========================================================
# SLIDE 6
# TECHNOLOGY + TESTING + FUTURE
# =========================================================

slide = prs.slides.add_slide(prs.slide_layouts[6])

add_background(slide)

add_header(
    slide,
    "Technology, Testing and Future Scope",
    "The current prototype and its path toward a larger civic platform",
    6
)


# Technology

add_card(
    slide,
    0.7,
    1.75,
    3.85,
    4.95,
    WHITE,
    LIGHT
)

add_text(
    slide,
    "TECHNOLOGY",
    1.0,
    2.05,
    3.25,
    0.4,
    size=17,
    color=PURPLE,
    bold=True,
    align=PP_ALIGN.CENTER
)

technology = [
    ("Frontend", "HTML, CSS, JavaScript"),
    ("Browser", "Geolocation and Speech APIs"),
    ("Backend", "Node.js and Express"),
    ("Upload", "Multer"),
    ("API", "REST API and CORS"),
    ("Deployment", "Vercel and Railway")
]

for i, (label, value) in enumerate(technology):

    y = 2.72 + i * 0.55

    add_text(
        slide,
        label,
        1.0,
        y,
        0.95,
        0.3,
        size=9,
        color=GREY,
        bold=True
    )

    add_text(
        slide,
        value,
        1.95,
        y,
        2.2,
        0.3,
        size=9.5,
        color=DARK,
        bold=True
    )


# Testing

add_card(
    slide,
    4.75,
    1.75,
    3.85,
    4.95,
    WHITE,
    LIGHT
)

add_text(
    slide,
    "TESTED WORKFLOWS",
    5.05,
    2.05,
    3.25,
    0.4,
    size=17,
    color=GREEN,
    bold=True,
    align=PP_ALIGN.CENTER
)

testing = [
    "Report submission",
    "Image upload",
    "Issue analysis",
    "Duplicate detection",
    "Admin workflow",
    "Status updates",
    "Citizen notifications",
    "Backend connection"
]

for i, item in enumerate(testing):

    y = 2.68 + i * 0.43

    add_card(
        slide,
        5.05,
        y,
        0.28,
        0.28,
        GREEN,
        GREEN
    )

    add_text(
        slide,
        "OK",
        5.05,
        y + 0.01,
        0.28,
        0.2,
        size=6.5,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER
    )

    add_text(
        slide,
        item,
        5.5,
        y - 0.02,
        2.65,
        0.3,
        size=10.5,
        color=DARK
    )


# Future

add_card(
    slide,
    8.8,
    1.75,
    3.85,
    4.95,
    WHITE,
    LIGHT
)

add_text(
    slide,
    "FUTURE SCOPE",
    9.1,
    2.05,
    3.25,
    0.4,
    size=17,
    color=BLUE,
    bold=True,
    align=PP_ALIGN.CENTER
)

future = [
    "Trained ML models",
    "Computer vision",
    "Cloud database",
    "Advanced GIS mapping",
    "Authority integration",
    "Secure authentication",
    "Advanced analytics",
    "Mobile application"
]

for i, item in enumerate(future):

    y = 2.68 + i * 0.43

    add_card(
        slide,
        9.1,
        y,
        0.28,
        0.28,
        BLUE,
        BLUE
    )

    add_text(
        slide,
        "+",
        9.1,
        y - 0.01,
        0.28,
        0.2,
        size=10,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER
    )

    add_text(
        slide,
        item,
        9.55,
        y - 0.02,
        2.65,
        0.3,
        size=10.5,
        color=DARK
    )


add_footer(slide)


# =========================================================
# SLIDE 7
# CONCLUSION
# LIGHT BACKGROUND + READABLE TEXT
# =========================================================

slide = prs.slides.add_slide(prs.slide_layouts[6])

add_background(slide, BG)


# Decorative elements

circle = slide.shapes.add_shape(
    MSO_SHAPE.OVAL,
    Inches(9.4),
    Inches(-1.1),
    Inches(4.7),
    Inches(4.7)
)

circle.fill.solid()
circle.fill.fore_color.rgb = PURPLE_LIGHT
circle.line.fill.background()


circle2 = slide.shapes.add_shape(
    MSO_SHAPE.OVAL,
    Inches(-1.0),
    Inches(5.3),
    Inches(3.5),
    Inches(3.5)
)

circle2.fill.solid()
circle2.fill.fore_color.rgb = BLUE_LIGHT
circle2.line.fill.background()


# Main title

add_text(
    slide,
    "CIVICEYE AI",
    0.9,
    0.95,
    11.5,
    0.65,
    size=34,
    color=DARK,
    bold=True,
    align=PP_ALIGN.CENTER
)

add_accent_line(
    slide,
    5.45,
    1.75,
    2.4,
    PURPLE
)


# Main statement

add_text(
    slide,
    "Smarter civic issue reporting\nthrough AI-assisted technology",
    1.6,
    2.25,
    10.1,
    1.25,
    size=28,
    color=DARK,
    bold=True,
    align=PP_ALIGN.CENTER
)


# Main workflow

workflow = [
    "REPORT",
    "ANALYSE",
    "PRIORITISE",
    "CONNECT",
    "MONITOR"
]

for i, item in enumerate(workflow):

    x = 1.05 + i * 2.38

    add_card(
        slide,
        x,
        4.15,
        1.95,
        0.7,
        WHITE,
        LIGHT
    )

    add_text(
        slide,
        item,
        x,
        4.34,
        1.95,
        0.3,
        size=10,
        color=DARK,
        bold=True,
        align=PP_ALIGN.CENTER
    )

    if i < len(workflow) - 1:

        add_arrow(
            slide,
            x + 1.95,
            4.5,
            x + 2.3,
            4.5,
            PURPLE
        )


# Closing statement

add_text(
    slide,
    "A working prototype designed to make civic issue reporting more structured, visible and manageable.",
    1.35,
    5.25,
    10.65,
    0.6,
    size=14,
    color=GREY,
    align=PP_ALIGN.CENTER
)


# Competition text

add_text(
    slide,
    "Young Computer Scientist Competition (YCS) 2026",
    1.0,
    6.0,
    11.3,
    0.4,
    size=15,
    color=DARK,
    bold=True,
    align=PP_ALIGN.CENTER
)


add_text(
    slide,
    "AI-Powered Web Application  |  Sri Lanka",
    1.0,
    6.48,
    11.3,
    0.35,
    size=12,
    color=PURPLE,
    bold=True,
    align=PP_ALIGN.CENTER
)


# =========================================================
# SAVE
# =========================================================

output_file = "CivicEye_AI_Premium_Final.pptx"

prs.save(output_file)

print()
print("================================================")
print(" CIVICEYE AI PREMIUM PPT CREATED SUCCESSFULLY")
print("================================================")
print()
print(f"File: {output_file}")
print()